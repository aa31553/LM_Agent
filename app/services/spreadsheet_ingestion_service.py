import csv
from datetime import date, datetime
from pathlib import Path
from typing import Any

import duckdb
import polars as pl
from openpyxl import load_workbook

from app.core.config import settings
from app.core.constants import ErrorCode
from app.core.exceptions import APIError
from app.services.spreadsheet_analysis_service import SpreadsheetAnalysisService
from app.storage.workspace_storage import WorkspaceStorage


class SpreadsheetIngestionService:
    """Convert XLSX/CSV sources into queryable Parquet datasets and build a profile."""

    def __init__(self, storage: WorkspaceStorage | None = None) -> None:
        self.storage = storage or WorkspaceStorage()
        self.legacy = SpreadsheetAnalysisService()

    def profile_and_convert(
        self,
        *,
        workspace_id,
        file_id,
        file_path: str,
        file_type: str,
    ) -> dict[str, Any]:
        inspection = self.legacy.inspect(file_path, file_type)
        if file_type == "xlsx":
            datasets = self._convert_xlsx(workspace_id, file_id, file_path)
        else:
            datasets = [
                self._convert_csv(
                    workspace_id,
                    file_id,
                    file_path,
                    dataset_key="csv",
                    sheet_name="CSV",
                )
            ]
        inspection_by_name = {item["name"]: item for item in inspection}
        for dataset in datasets:
            inspected = inspection_by_name.get(dataset["sheet"])
            if inspected is not None:
                dataset["sample_rows"] = inspected.get("sample_rows", [])
                dataset["warnings"] = inspected.get("warnings", [])
        manifest = {
            "schema_version": "1.0",
            "conversion_engine": "duckdb",
            "query_engine": "polars",
            "datasets": datasets,
        }
        warnings = [warning for dataset in datasets for warning in dataset.get("warnings", [])]
        profile = {
            "schema_version": "2.0",
            "sheets": inspection,
            "datasets": datasets,
            "warnings": warnings,
            "profiled_at": datetime.utcnow().isoformat(),
        }
        profile_path = self.storage.save_profile(
            workspace_id,
            file_id,
            profile,
        )
        return {
            "profile_path": profile_path,
            "profile": profile,
            "dataset_manifest": manifest,
        }

    def _convert_xlsx(
        self,
        workspace_id,
        file_id,
        file_path: str,
    ) -> list[dict[str, Any]]:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        datasets: list[dict[str, Any]] = []
        try:
            for index, worksheet in enumerate(workbook.worksheets, start=1):
                dataset_key = f"sheet_{index}"
                parquet_path = self.storage.dataset_path(
                    workspace_id,
                    file_id,
                    dataset_key,
                )
                temporary_csv = parquet_path.with_suffix(".source.csv")
                temporary_csv.parent.mkdir(parents=True, exist_ok=True)
                iterator = worksheet.iter_rows(values_only=True)
                try:
                    raw_headers = next(iterator)
                except StopIteration:
                    raw_headers = ()
                headers = self.legacy._normalize_headers(raw_headers)
                if len(headers) > settings.analysis_profile_max_columns:
                    raise APIError(
                        ErrorCode.INVALID_REQUEST,
                        "Worksheet exceeds the configured profiling column limit.",
                        413,
                        details={
                            "sheet": worksheet.title,
                            "analysis_profile_max_columns": (settings.analysis_profile_max_columns),
                        },
                    )
                with temporary_csv.open("w", encoding="utf-8", newline="") as handle:
                    writer = csv.writer(handle)
                    writer.writerow(headers)
                    for values in iterator:
                        writer.writerow([self._csv_value(value) for value in values])
                try:
                    self._csv_to_parquet(temporary_csv, parquet_path)
                finally:
                    temporary_csv.unlink(missing_ok=True)
                datasets.append(
                    self._profile_parquet(
                        dataset_key=dataset_key,
                        sheet_name=worksheet.title,
                        parquet_path=parquet_path,
                    )
                )
        finally:
            workbook.close()
        return datasets

    def _convert_csv(
        self,
        workspace_id,
        file_id,
        file_path: str,
        *,
        dataset_key: str,
        sheet_name: str,
    ) -> dict[str, Any]:
        parquet_path = self.storage.dataset_path(workspace_id, file_id, dataset_key)
        source = Path(file_path)
        encoding = self.legacy._detect_encoding(file_path)
        temporary_csv: Path | None = None
        if encoding not in {"utf-8", "utf-8-sig"}:
            temporary_csv = parquet_path.with_suffix(".source.csv")
            temporary_csv.parent.mkdir(parents=True, exist_ok=True)
            with (
                source.open("r", encoding=encoding, newline="") as reader,
                temporary_csv.open("w", encoding="utf-8", newline="") as writer,
            ):
                for chunk in iter(lambda: reader.read(1024 * 1024), ""):
                    writer.write(chunk)
            source = temporary_csv
        try:
            self._csv_to_parquet(source, parquet_path)
        finally:
            if temporary_csv is not None:
                temporary_csv.unlink(missing_ok=True)
        return self._profile_parquet(
            dataset_key=dataset_key,
            sheet_name=sheet_name,
            parquet_path=parquet_path,
        )

    def _csv_to_parquet(self, source: Path, destination: Path) -> None:
        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination.with_suffix(".parquet.tmp")
        connection = duckdb.connect()
        try:
            source_sql = self._sql_literal(str(source.resolve()))
            destination_sql = self._sql_literal(str(temporary.resolve()))
            connection.execute(
                "COPY (SELECT * FROM read_csv_auto("
                f"{source_sql}, header=true, sample_size=-1, all_varchar=false, "
                "ignore_errors=false, normalize_names=false"
                f")) TO {destination_sql} (FORMAT PARQUET, COMPRESSION ZSTD)"
            )
        finally:
            connection.close()
        temporary.replace(destination)

    def _profile_parquet(
        self,
        *,
        dataset_key: str,
        sheet_name: str,
        parquet_path: Path,
    ) -> dict[str, Any]:
        lazy = pl.scan_parquet(parquet_path)
        schema = lazy.collect_schema()
        if len(schema) > settings.analysis_profile_max_columns:
            raise APIError(
                ErrorCode.INVALID_REQUEST,
                "Dataset exceeds the configured profiling column limit.",
                413,
            )
        row_count = int(lazy.select(pl.len()).collect().item())
        null_values = (
            lazy.select(pl.all().null_count()).collect().row(0, named=True) if schema else {}
        )
        samples = (
            lazy.head(settings.analysis_profile_sample_rows).collect().to_dicts() if schema else []
        )
        columns = [
            {
                "name": name,
                "inferred_type": self._polars_type(dtype),
                "physical_type": str(dtype),
                "null_count": int(null_values.get(name, 0)),
                "sample_values": [
                    self._json_value(row.get(name)) for row in samples if row.get(name) is not None
                ][:5],
            }
            for name, dtype in schema.items()
        ]
        return {
            "key": dataset_key,
            "sheet": sheet_name,
            "path": str(parquet_path),
            "row_count": row_count,
            "column_count": len(columns),
            "columns": columns,
        }

    @staticmethod
    def _polars_type(dtype: Any) -> str:
        if dtype.is_integer():
            return "integer"
        if dtype.is_float() or dtype.is_decimal():
            return "number"
        if dtype.is_temporal():
            return "datetime"
        if dtype == pl.Boolean:
            return "boolean"
        return "string"

    @staticmethod
    def _csv_value(value: Any) -> Any:
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        return value

    @staticmethod
    def _json_value(value: Any) -> Any:
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        return value

    @staticmethod
    def _sql_literal(value: str) -> str:
        return "'" + value.replace("'", "''") + "'"


def profile_spreadsheet_file(
    *,
    workspace_id,
    file_id,
    file_path: str,
    file_type: str,
    storage_root: str | None = None,
) -> dict[str, Any]:
    storage = WorkspaceStorage(root=storage_root) if storage_root is not None else None
    return SpreadsheetIngestionService(storage).profile_and_convert(
        workspace_id=workspace_id,
        file_id=file_id,
        file_path=file_path,
        file_type=file_type,
    )
