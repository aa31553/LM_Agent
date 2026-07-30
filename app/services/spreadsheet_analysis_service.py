import csv
import math
from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, ClassVar

from openpyxl import load_workbook

from app.core.config import settings
from app.core.constants import ErrorCode
from app.core.exceptions import APIError
from app.schemas.analysis import AnalysisPlan, FilterCondition
from app.services.office_file_preparation_service import OfficeFilePreparationService


class SpreadsheetAnalysisService:
    """Deterministic, streaming spreadsheet inspection and aggregation.

    The service intentionally does not execute LLM-generated Python or SQL. Every
    supported operation is represented by a validated AnalysisPlan whitelist.
    """

    SUPPORTED_TYPES: ClassVar[set[str]] = {"xlsx", "csv"}

    def __init__(
        self,
        office_preparation: OfficeFilePreparationService | None = None,
    ) -> None:
        self.office_preparation = office_preparation or OfficeFilePreparationService()

    def inspect(self, file_path: str, file_type: str) -> list[dict[str, Any]]:
        normalized_type = self._ensure_supported(file_type)
        if normalized_type == "xlsx":
            with self.office_preparation.prepare_bounded_sync(
                file_path, normalized_type
            ) as prepared:
                return self.inspect_prepared(str(prepared.path), normalized_type)
        return self.inspect_prepared(file_path, normalized_type)

    def inspect_prepared(
        self,
        file_path: str,
        file_type: str,
    ) -> list[dict[str, Any]]:
        normalized_type = self._ensure_supported(file_type)
        if normalized_type == "xlsx":
            return self._inspect_xlsx(file_path)
        return [self._inspect_csv(file_path)]

    def validate_plan(
        self,
        file_path: str,
        file_type: str,
        plan: AnalysisPlan,
    ) -> tuple[AnalysisPlan, list[str]]:
        normalized_type = self._ensure_supported(file_type)
        if normalized_type == "xlsx":
            with self.office_preparation.prepare_bounded_sync(
                file_path, normalized_type
            ) as prepared:
                return self.validate_prepared(
                    str(prepared.path),
                    normalized_type,
                    plan,
                )
        return self.validate_prepared(file_path, normalized_type, plan)

    def validate_prepared(
        self,
        file_path: str,
        file_type: str,
        plan: AnalysisPlan,
    ) -> tuple[AnalysisPlan, list[str]]:
        if (
            plan.sources
            or plan.joins
            or plan.date_buckets
            or plan.pivot is not None
            or plan.correlation is not None
            or any(item.function in {"distinct_count", "percentile"} for item in plan.aggregations)
        ):
            raise APIError(
                ErrorCode.INVALID_REQUEST,
                "Advanced operations require preprocessed dataset sources.",
                400,
            )
        sheets = self.inspect_prepared(file_path, file_type)
        selected_sheet = self._select_sheet(sheets, plan.sheet)
        source_headers = {item["name"]: item for item in selected_sheet["columns"]}
        aggregation_aliases = {item.alias for item in plan.aggregations}

        source_references = set(plan.select) | set(plan.group_by)
        source_references.update(condition.column for condition in plan.filters)
        source_references.update(item.column for item in plan.aggregations if item.column)
        source_references.update(
            item.condition.column for item in plan.aggregations if item.condition is not None
        )
        missing_source = sorted(name for name in source_references if name not in source_headers)
        if missing_source:
            raise APIError(
                ErrorCode.INVALID_REQUEST,
                "Analysis plan references source columns that do not exist.",
                400,
                details={"missing_columns": missing_source},
            )

        output_columns = (
            set(plan.group_by) | aggregation_aliases if plan.aggregations else set(plan.select)
        )
        output_references = {item.column for item in plan.sort}
        output_references.update(item.x_field for item in plan.charts)
        output_references.update(item.y_field for item in plan.charts)
        output_references.update(
            item.series_field for item in plan.charts if item.series_field is not None
        )
        output_references.update(field for item in plan.charts for field in item.tooltip_fields)
        missing_output = sorted(name for name in output_references if name not in output_columns)
        if missing_output:
            raise APIError(
                ErrorCode.INVALID_REQUEST,
                "Sort or chart fields are not present in the analysis output.",
                400,
                details={"missing_output_fields": missing_output},
            )

        for aggregation in plan.aggregations:
            if aggregation.function in {"sum", "mean", "std"} and aggregation.column:
                inferred = source_headers[aggregation.column]["inferred_type"]
                if inferred not in {"integer", "number", "unknown"}:
                    raise APIError(
                        ErrorCode.INVALID_REQUEST,
                        f"Column {aggregation.column} is not numeric.",
                        400,
                    )

        if plan.sort and not plan.aggregations:
            raise APIError(
                ErrorCode.INVALID_REQUEST,
                "Sorting raw extracted rows is not supported in the first "
                "deterministic analysis version; add an aggregation or remove sort.",
                400,
            )
        normalized = plan.model_copy(update={"sheet": selected_sheet["name"]})
        warnings = list(selected_sheet.get("warnings", []))
        if normalized.select and normalized.aggregations:
            warnings.append(
                "select is ignored for aggregation output; use group_by and aggregation aliases."
            )
        if normalized.aggregations and not normalized.group_by:
            warnings.append("The aggregation produces one overall result row.")
        return normalized, warnings

    def execute(
        self,
        file_path: str,
        file_type: str,
        plan: AnalysisPlan,
    ) -> dict[str, Any]:
        normalized_type = self._ensure_supported(file_type)
        if normalized_type == "xlsx":
            with self.office_preparation.prepare_bounded_sync(
                file_path, normalized_type
            ) as prepared:
                return self.execute_prepared(
                    str(prepared.path),
                    normalized_type,
                    plan,
                )
        return self.execute_prepared(file_path, normalized_type, plan)

    def execute_prepared(
        self,
        file_path: str,
        file_type: str,
        plan: AnalysisPlan,
    ) -> dict[str, Any]:
        normalized, warnings = self.validate_prepared(file_path, file_type, plan)
        processed_rows = 0
        matched_rows = 0
        extracted_rows: list[dict[str, Any]] = []
        groups: dict[tuple[Any, ...], dict[str, Any]] = {}

        for row in self._iter_rows(file_path, file_type, normalized.sheet):
            processed_rows += 1
            if processed_rows > settings.analysis_max_rows:
                raise APIError(
                    ErrorCode.INVALID_REQUEST,
                    "Spreadsheet exceeds the configured analysis row limit.",
                    413,
                    details={"analysis_max_rows": settings.analysis_max_rows},
                )
            if not all(self._matches(row.get(item.column), item) for item in normalized.filters):
                continue
            matched_rows += 1

            if not normalized.aggregations:
                if len(extracted_rows) < min(
                    normalized.limit,
                    settings.analysis_result_max_rows,
                ):
                    extracted_rows.append(
                        {name: self._json_value(row.get(name)) for name in normalized.select}
                    )
                continue

            key = tuple(self._hashable(row.get(name)) for name in normalized.group_by)
            if key not in groups:
                if len(groups) >= settings.analysis_max_groups:
                    raise APIError(
                        ErrorCode.INVALID_REQUEST,
                        "Analysis produced too many groups.",
                        413,
                        details={"analysis_max_groups": settings.analysis_max_groups},
                    )
                groups[key] = {item.alias: self._new_state() for item in normalized.aggregations}
            for item in normalized.aggregations:
                self._update_state(
                    groups[key][item.alias],
                    item.function,
                    row,
                    item.column,
                    item.condition,
                )

        if normalized.aggregations:
            result_rows: list[dict[str, Any]] = []
            for key, states in groups.items():
                output = {
                    name: self._json_value(value)
                    for name, value in zip(normalized.group_by, key, strict=True)
                }
                for item in normalized.aggregations:
                    output[item.alias] = self._json_value(
                        self._finalize_state(states[item.alias], item.function)
                    )
                result_rows.append(output)
        else:
            result_rows = extracted_rows

        for sort in reversed(normalized.sort):
            result_rows.sort(
                key=lambda item, column=sort.column: self._sort_key(item.get(column)),
                reverse=sort.direction == "desc",
            )

        total_result_rows = len(result_rows) if normalized.aggregations else matched_rows
        result_rows = result_rows[: min(normalized.limit, settings.analysis_result_max_rows)]
        columns = (
            list(result_rows[0])
            if result_rows
            else (
                normalized.group_by + [item.alias for item in normalized.aggregations]
                if normalized.aggregations
                else normalized.select
            )
        )
        charts = [self._chart_payload(chart, result_rows) for chart in normalized.charts]
        return {
            "summary": {
                "processed_rows": processed_rows,
                "matched_rows": matched_rows,
                "result_rows": total_result_rows,
                "returned_rows": len(result_rows),
                "truncated": total_result_rows > len(result_rows),
                "warnings": warnings,
            },
            "table": {
                "columns": [{"key": name, "label": name} for name in columns],
                "rows": result_rows,
            },
            "charts": charts,
            "plan": normalized.model_dump(mode="json"),
        }

    def _inspect_xlsx(self, file_path: str) -> list[dict[str, Any]]:
        value_workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )
        formula_workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=False,
        )
        try:
            return [
                self._inspect_worksheet(
                    value_sheet,
                    formula_workbook[value_sheet.title],
                )
                for value_sheet in value_workbook.worksheets
            ]
        finally:
            value_workbook.close()
            formula_workbook.close()

    def _inspect_worksheet(self, value_worksheet, formula_worksheet) -> dict[str, Any]:
        value_iterator = value_worksheet.iter_rows()
        formula_iterator = formula_worksheet.iter_rows()
        try:
            header_cells = next(value_iterator)
            headers = self._normalize_headers([cell.value for cell in header_cells])
            next(formula_iterator, ())
        except StopIteration:
            headers = []
        samples: list[dict[str, Any]] = []
        formula_cells = 0
        formulas_without_cached_values = 0
        formatted_cells = 0
        for value_cells in value_iterator:
            source_cells = next(formula_iterator, ())
            values = [cell.value for cell in value_cells]
            samples.append(
                {
                    header: self._json_value(value)
                    for header, value in zip(headers, values, strict=False)
                }
            )
            for value_cell, source_cell in zip(
                value_cells,
                source_cells,
                strict=False,
            ):
                if source_cell.data_type == "f":
                    formula_cells += 1
                    if value_cell.value is None:
                        formulas_without_cached_values += 1
                if (
                    source_cell.value not in {None, ""}
                    and source_cell.number_format
                    and source_cell.number_format != "General"
                ):
                    formatted_cells += 1
            if len(samples) >= settings.analysis_sample_rows:
                break
        warnings = self._xlsx_sample_warnings(
            value_worksheet.title,
            formula_cells=formula_cells,
            formulas_without_cached_values=formulas_without_cached_values,
            formatted_cells=formatted_cells,
        )
        return self._sheet_info(
            value_worksheet.title,
            max((value_worksheet.max_row or 1) - 1, 0),
            headers,
            samples,
            formula_cells_in_sample=formula_cells,
            formatted_cells_in_sample=formatted_cells,
            warnings=warnings,
        )

    def _inspect_csv(self, file_path: str) -> dict[str, Any]:
        encoding = self._detect_encoding(file_path)
        with open(file_path, "r", encoding=encoding, newline="") as handle:
            reader = csv.reader(handle)
            try:
                headers = self._normalize_headers(next(reader))
            except StopIteration:
                headers = []
            samples: list[dict[str, Any]] = []
            for values in reader:
                samples.append(
                    {
                        header: self._json_value(value)
                        for header, value in zip(headers, values, strict=False)
                    }
                )
                if len(samples) >= settings.analysis_sample_rows:
                    break
        return self._sheet_info("CSV", None, headers, samples)

    def _sheet_info(
        self,
        name: str,
        row_count: int | None,
        headers: list[str],
        samples: list[dict[str, Any]],
        *,
        formula_cells_in_sample: int = 0,
        formatted_cells_in_sample: int = 0,
        warnings: list[str] | None = None,
    ) -> dict[str, Any]:
        columns = []
        for header in headers:
            values = [row.get(header) for row in samples]
            columns.append(
                {
                    "name": header,
                    "inferred_type": self._infer_type(values),
                    "sample_values": [value for value in values if value not in {None, ""}][:5],
                    "null_count_in_sample": sum(value is None or value == "" for value in values),
                }
            )
        return {
            "name": name,
            "row_count": row_count,
            "column_count": len(headers),
            "columns": columns,
            "sample_rows": samples,
            "formula_cells_in_sample": formula_cells_in_sample,
            "formatted_cells_in_sample": formatted_cells_in_sample,
            "warnings": warnings or [],
        }

    def _xlsx_sample_warnings(
        self,
        sheet_name: str,
        *,
        formula_cells: int,
        formulas_without_cached_values: int,
        formatted_cells: int,
    ) -> list[str]:
        warnings: list[str] = []
        if formula_cells:
            warnings.append(
                f"Sheet {sheet_name}: {formula_cells} formula cell(s) were found "
                "in the inspected sample. Analysis uses the last cached values "
                "saved by Excel and does not recalculate formulas."
            )
        if formulas_without_cached_values:
            warnings.append(
                f"Sheet {sheet_name}: {formulas_without_cached_values} sampled "
                "formula cell(s) have no cached value. Recalculate and save the "
                "workbook in Excel before analysis."
            )
        if formatted_cells:
            warnings.append(
                f"Sheet {sheet_name}: {formatted_cells} sampled cell(s) use Excel "
                "display formatting. Analysis uses underlying values; number, "
                "date, percentage, and leading-zero display formats are not "
                "preserved in result JSON."
            )
        return warnings

    def _iter_rows(
        self,
        file_path: str,
        file_type: str,
        sheet_name: str | None,
    ) -> Iterator[dict[str, Any]]:
        normalized_type = self._ensure_supported(file_type)
        if normalized_type == "xlsx":
            yield from self._iter_xlsx_rows(file_path, sheet_name)
        else:
            yield from self._iter_csv_rows(file_path)

    def _iter_xlsx_rows(
        self,
        file_path: str,
        sheet_name: str | None,
    ) -> Iterator[dict[str, Any]]:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise APIError(
                    ErrorCode.INVALID_REQUEST,
                    "Worksheet not found.",
                    404,
                )
            worksheet = workbook[sheet_name]
            iterator = worksheet.iter_rows(values_only=True)
            try:
                headers = self._normalize_headers(next(iterator))
            except StopIteration:
                return
            for values in iterator:
                yield dict(zip(headers, values, strict=False))
        finally:
            workbook.close()

    def _iter_csv_rows(self, file_path: str) -> Iterator[dict[str, Any]]:
        encoding = self._detect_encoding(file_path)
        with open(file_path, "r", encoding=encoding, newline="") as handle:
            reader = csv.reader(handle)
            try:
                headers = self._normalize_headers(next(reader))
            except StopIteration:
                return
            for values in reader:
                yield dict(zip(headers, values, strict=False))

    def _select_sheet(
        self,
        sheets: list[dict[str, Any]],
        requested: str | None,
    ) -> dict[str, Any]:
        if not sheets:
            raise APIError(
                ErrorCode.INVALID_REQUEST,
                "Spreadsheet has no worksheets.",
                400,
            )
        if requested is None:
            return sheets[0]
        for sheet in sheets:
            if sheet["name"] == requested:
                return sheet
        raise APIError(
            ErrorCode.INVALID_REQUEST,
            "Worksheet not found.",
            404,
            details={"available_sheets": [sheet["name"] for sheet in sheets]},
        )

    def _normalize_headers(self, values) -> list[str]:
        result: list[str] = []
        counts: dict[str, int] = {}
        for index, value in enumerate(values or (), start=1):
            base = str(value).strip() if value not in {None, ""} else f"column_{index}"
            counts[base] = counts.get(base, 0) + 1
            result.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
        return result

    def _infer_type(self, values: list[Any]) -> str:
        non_empty = [value for value in values if value is not None and value != ""]
        if not non_empty:
            return "unknown"
        if all(isinstance(value, bool) for value in non_empty):
            return "boolean"
        if all(isinstance(value, (date, datetime)) for value in non_empty):
            return "datetime"
        if all(self._is_integer(value) for value in non_empty):
            return "integer"
        if all(self._to_number(value) is not None for value in non_empty):
            return "number"
        return "string"

    def _matches(self, value: Any, condition: FilterCondition) -> bool:
        operator = condition.operator
        target = condition.value
        if operator == "is_null":
            return value is None or value == ""
        if operator == "not_null":
            return value is not None and value != ""
        if operator == "contains":
            return str(target).lower() in str(value or "").lower()
        if operator == "in":
            candidates = target if isinstance(target, list) else [target]
            return value in candidates or str(value) in {str(item) for item in candidates}

        left_number = self._to_number(value)
        right_number = self._to_number(target)
        left, right = (
            (left_number, right_number)
            if left_number is not None and right_number is not None
            else (str(value), str(target))
        )
        return {
            "eq": left == right,
            "ne": left != right,
            "gt": left > right,
            "gte": left >= right,
            "lt": left < right,
            "lte": left <= right,
        }[operator]

    def _new_state(self) -> dict[str, Any]:
        return {
            "count": 0,
            "sum": 0.0,
            "mean": 0.0,
            "m2": 0.0,
            "min": None,
            "max": None,
        }

    def _update_state(
        self,
        state: dict[str, Any],
        function: str,
        row: dict[str, Any],
        column: str | None,
        condition: FilterCondition | None,
    ) -> None:
        if function == "count_if":
            if condition is not None and self._matches(
                row.get(condition.column),
                condition,
            ):
                state["count"] += 1
            return
        if function == "count":
            if column is None or row.get(column) not in {None, ""}:
                state["count"] += 1
            return

        value = row.get(column) if column else None
        if value in {None, ""}:
            return
        if function in {"sum", "mean", "std"}:
            numeric = self._to_number(value)
            if numeric is None:
                return
            state["count"] += 1
            state["sum"] += numeric
            delta = numeric - state["mean"]
            state["mean"] += delta / state["count"]
            state["m2"] += delta * (numeric - state["mean"])
            return
        state["count"] += 1
        if state["min"] is None or self._safe_compare(
            value,
            state["min"],
            "lt",
        ):
            state["min"] = value
        if state["max"] is None or self._safe_compare(
            value,
            state["max"],
            "gt",
        ):
            state["max"] = value

    def _finalize_state(self, state: dict[str, Any], function: str) -> Any:
        if function in {"count", "count_if"}:
            return state["count"]
        if function == "sum":
            return state["sum"]
        if function == "mean":
            return state["mean"] if state["count"] else None
        if function == "std":
            return math.sqrt(state["m2"] / (state["count"] - 1)) if state["count"] > 1 else 0.0
        if function == "min":
            return state["min"]
        if function == "max":
            return state["max"]
        raise APIError(
            ErrorCode.INVALID_REQUEST,
            f"Unsupported aggregation: {function}",
            400,
        )

    def _chart_payload(
        self,
        chart,
        rows: list[dict[str, Any]],
    ) -> dict[str, Any]:
        return {
            "schema_version": "2.0",
            "chart_type": chart.type,
            "type": chart.type,
            "title": chart.title,
            "x_field": chart.x_field,
            "y_field": chart.y_field,
            "series_field": chart.series_field,
            "tooltip_fields": chart.tooltip_fields,
            "x_type": chart.x_type,
            "format": {
                "y_unit": chart.y_unit,
                "decimal_places": chart.decimal_places,
            },
            "interaction": {
                "tooltip": True,
                "zoom": chart.zoom,
                "export": True,
            },
            "data": [
                {
                    field: self._json_value(row.get(field))
                    for field in dict.fromkeys(
                        [
                            chart.x_field,
                            chart.y_field,
                            chart.series_field,
                            *chart.tooltip_fields,
                        ]
                    )
                    if field is not None
                }
                for row in rows
            ],
        }

    def _detect_encoding(self, file_path: str) -> str:
        sample = Path(file_path).read_bytes()[:65536]
        for encoding in ("utf-8-sig", "utf-8", "cp950", "big5"):
            try:
                sample.decode(encoding)
                return encoding
            except UnicodeDecodeError:
                continue
        raise APIError(
            ErrorCode.INVALID_REQUEST,
            "CSV encoding is not supported; use UTF-8 or Big5/CP950.",
            400,
        )

    def _ensure_supported(self, file_type: str) -> str:
        normalized = file_type.lower().lstrip(".")
        if normalized not in self.SUPPORTED_TYPES:
            raise APIError(
                ErrorCode.INVALID_REQUEST,
                "The analysis workspace accepts only .xlsx and .csv files.",
                400,
            )
        return normalized

    def _to_number(self, value: Any) -> float | None:
        if value is None or value == "" or isinstance(value, bool):
            return None
        try:
            return float(value)
        except (TypeError, ValueError):
            return None

    def _is_integer(self, value: Any) -> bool:
        if isinstance(value, bool):
            return False
        if isinstance(value, int):
            return True
        if isinstance(value, float):
            return value.is_integer()
        if isinstance(value, str):
            try:
                return float(value).is_integer()
            except ValueError:
                return False
        return False

    def _sort_key(self, value: Any) -> tuple[int, Any]:
        if value is None:
            return 3, ""
        number = self._to_number(value)
        if number is not None:
            return 0, number
        if isinstance(value, (date, datetime)):
            return 1, value.isoformat()
        return 2, str(value)

    def _safe_compare(self, left: Any, right: Any, operator: str) -> bool:
        left_number = self._to_number(left)
        right_number = self._to_number(right)
        if left_number is not None and right_number is not None:
            left_value, right_value = left_number, right_number
        else:
            left_value, right_value = str(left), str(right)
        return left_value < right_value if operator == "lt" else left_value > right_value

    def _json_value(self, value: Any) -> Any:
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        if isinstance(value, Decimal):
            return float(value)
        return str(value)

    def _hashable(self, value: Any) -> Any:
        try:
            hash(value)
            return value
        except TypeError:
            return self._json_value(value)


def execute_spreadsheet_analysis(
    file_path: str,
    file_type: str,
    plan_payload: dict[str, Any],
) -> dict[str, Any]:
    """Pickle-safe process entrypoint used by the analysis worker."""

    return SpreadsheetAnalysisService().execute(
        file_path,
        file_type,
        AnalysisPlan.model_validate(plan_payload),
    )
