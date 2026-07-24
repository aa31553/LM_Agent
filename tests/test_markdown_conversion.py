from pathlib import Path
import zipfile

import pytest
from openpyxl import Workbook
from pptx import Presentation

from app.services.chunking_service import ChunkingService
from app.services.markdown_conversion_service import MarkdownConversionService
from app.storage.local_storage import LocalStorage


class FakeMarkItDown:
    def __init__(self) -> None:
        self.paths: list[Path] = []

    def convert_local(self, path: Path):
        self.paths.append(path)
        return type(
            "ConversionResult",
            (),
            {"text_content": "# Report\r\n\r\n| Metric | Value |\r\n| --- | --- |"},
        )()


@pytest.mark.asyncio
async def test_conversion_uses_local_only_api_and_normalizes_markdown(tmp_path: Path) -> None:
    source = tmp_path / "report.pdf"
    source.write_bytes(b"test")
    converter = FakeMarkItDown()

    markdown = await MarkdownConversionService(converter).convert_local(source)

    assert converter.paths == [source]
    assert markdown == "# Report\n\n| Metric | Value |\n| --- | --- |\n"


@pytest.mark.asyncio
async def test_local_storage_separates_originals_and_markdown(tmp_path: Path) -> None:
    storage = LocalStorage(str(tmp_path / "uploads"))

    original_path = Path(await storage.save_original("document.pdf", b"pdf"))
    markdown_path = Path(await storage.save_markdown("document.md", "# Document\n"))

    assert original_path == tmp_path / "uploads" / "originals" / "document.pdf"
    assert markdown_path == tmp_path / "uploads" / "markdown" / "document.md"
    assert original_path.read_bytes() == b"pdf"
    assert markdown_path.read_text(encoding="utf-8") == "# Document\n"


def test_markdown_chunking_preserves_structure_and_marks_provenance() -> None:
    markdown = "# Heading\n\n| A | B |\n| --- | --- |\n| one | two |\n"

    chunks = ChunkingService().chunk_markdown(markdown, markdown_path="markdown/report.md")

    assert len(chunks) == 1
    assert chunks[0].content == markdown.strip()
    assert chunks[0].metadata == {
        "source_type": "markdown",
        "markdown_path": "markdown/report.md",
    }


@pytest.mark.asyncio
async def test_real_markitdown_converts_supported_office_formats(tmp_path: Path) -> None:
    docx_path = tmp_path / "policy.docx"
    with zipfile.ZipFile(docx_path, "w") as archive:
        archive.writestr(
            "word/document.xml",
            """
            <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
              <w:body><w:p><w:r><w:t>Retention policy</w:t></w:r></w:p></w:body>
            </w:document>
            """,
        )

    xlsx_path = tmp_path / "risks.xlsx"
    workbook = Workbook()
    workbook.active.append(["Metric", "Value"])
    workbook.active.append(["Data leakage", 8])
    workbook.save(xlsx_path)

    pptx_path = tmp_path / "roadmap.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Security roadmap"
    slide.placeholders[1].text = "Enable agent tools carefully."
    deck.save(pptx_path)

    converter = MarkdownConversionService()
    converted = {
        path.suffix: await converter.convert_local(path)
        for path in (docx_path, xlsx_path, pptx_path)
    }

    assert "Retention policy" in converted[".docx"]
    assert "| Metric | Value |" in converted[".xlsx"]
    assert "# Security roadmap" in converted[".pptx"]
