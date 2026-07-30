import json
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.core.constants import AnalysisFileStatus, ConfidentialLevel
from app.core.exceptions import APIError
from app.core.security import Principal
from app.db.base import Base
from app.models.analysis import AnalysisFile, AnalysisJob
from app.models.chat import ChatSession
from app.schemas.analysis import (
    AggregationSpec,
    AnalysisChartArtifactRequest,
    AnalysisExportRequest,
    AnalysisHybridRequest,
    AnalysisPlan,
    AnalysisPlanDraftCreate,
    AnalysisReportRequest,
    ChartSpec,
    CorrelationSpec,
    DatasetSource,
    DateBucketSpec,
    JoinSpec,
    PivotSpec,
)
from app.services.analysis_artifact_service import AnalysisArtifactService
from app.services.analysis_orchestrator_service import AnalysisOrchestratorService
from app.services.audit_service import AuditService
from app.services.dataset_query_service import DatasetQueryService
from app.services.hybrid_analysis_service import HybridAnalysisService
from app.services.spreadsheet_ingestion_service import SpreadsheetIngestionService
from app.services.vector_store_service import RetrievedChunk
from app.services.workspace_service import WorkspaceService
from app.storage.workspace_storage import WorkspaceStorage


def _principal() -> Principal:
    return Principal(
        external_user_id="phase34-user",
        username="phase34-user",
        clearance_level=ConfidentialLevel.INTERNAL,
    )


def _workbook(path: Path, rows: list[list], *, sheet: str = "Data") -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


def test_background_preprocessing_converts_each_sheet_to_parquet(
    tmp_path: Path,
) -> None:
    path = tmp_path / "production.xlsx"
    workbook = Workbook()
    production = workbook.active
    production.title = "Production"
    production.append(["Date", "Machine", "Yield"])
    production.append(["2026-07-01", "A", 98])
    machines = workbook.create_sheet("Machines")
    machines.append(["Machine", "Line"])
    machines.append(["A", "L1"])
    workbook.save(path)

    workspace_id = uuid4()
    file_id = uuid4()
    result = SpreadsheetIngestionService(
        WorkspaceStorage(root=str(tmp_path / "storage"))
    ).profile_and_convert(
        workspace_id=workspace_id,
        file_id=file_id,
        file_path=str(path),
        file_type="xlsx",
    )

    manifest = result["dataset_manifest"]
    assert manifest["conversion_engine"] == "duckdb"
    assert manifest["query_engine"] == "polars"
    assert [item["sheet"] for item in manifest["datasets"]] == [
        "Production",
        "Machines",
    ]
    assert all(Path(item["path"]).is_file() for item in manifest["datasets"])
    assert manifest["datasets"][0]["row_count"] == 1
    assert result["profile"]["schema_version"] == "2.0"


def test_polars_query_supports_join_date_bucket_percentile_pivot_and_correlation(
    tmp_path: Path,
) -> None:
    storage = WorkspaceStorage(root=str(tmp_path / "storage"))
    workspace_id = uuid4()
    production_file_id = uuid4()
    machine_file_id = uuid4()
    production_path = tmp_path / "production.xlsx"
    machine_path = tmp_path / "machines.xlsx"
    _workbook(
        production_path,
        [
            ["Date", "Machine", "Yield", "Pressure"],
            ["2026-07-01", "A", 98, 10],
            ["2026-07-15", "A", 96, 12],
            ["2026-08-01", "B", 90, 20],
        ],
        sheet="Production",
    )
    _workbook(
        machine_path,
        [["Machine", "Line"], ["A", "L1"], ["B", "L2"]],
        sheet="Machines",
    )
    production = SpreadsheetIngestionService(storage).profile_and_convert(
        workspace_id=workspace_id,
        file_id=production_file_id,
        file_path=str(production_path),
        file_type="xlsx",
    )["dataset_manifest"]
    machines = SpreadsheetIngestionService(storage).profile_and_convert(
        workspace_id=workspace_id,
        file_id=machine_file_id,
        file_path=str(machine_path),
        file_type="xlsx",
    )["dataset_manifest"]
    manifests = {
        str(production_file_id): production,
        str(machine_file_id): machines,
    }
    with pytest.raises(APIError, match="connected by an ordered join"):
        DatasetQueryService().validate_plan(
            AnalysisPlan(
                sources=[
                    DatasetSource(
                        file_id=production_file_id,
                        alias="p",
                        sheet="Production",
                    ),
                    DatasetSource(
                        file_id=machine_file_id,
                        alias="m",
                        sheet="Machines",
                    ),
                ],
                select=["p.Machine"],
            ),
            manifests,
        )
    plan = AnalysisPlan(
        sources=[
            DatasetSource(
                file_id=production_file_id,
                alias="p",
                sheet="Production",
            ),
            DatasetSource(
                file_id=machine_file_id,
                alias="m",
                sheet="Machines",
            ),
        ],
        joins=[
            JoinSpec(
                left_alias="p",
                right_alias="m",
                left_on=["Machine"],
                right_on=["Machine"],
            )
        ],
        date_buckets=[DateBucketSpec(column="p.Date", unit="month", alias="month")],
        group_by=["month", "m.Line"],
        aggregations=[
            AggregationSpec(
                function="percentile",
                column="p.Yield",
                alias="median_yield",
                percentile=0.5,
            )
        ],
        charts=[
            ChartSpec(
                type="line",
                x_field="month",
                y_field="median_yield",
                series_field="m.Line",
                x_type="time",
                y_unit="%",
            )
        ],
    )
    result = DatasetQueryService().execute(plan, manifests)
    assert result["summary"]["query_engine"] == "polars"
    assert {row["median_yield"] for row in result["table"]["rows"]} == {90, 97}
    assert result["charts"][0]["schema_version"] == "2.0"
    assert result["charts"][0]["series_field"] == "m.Line"

    pivot = DatasetQueryService().execute(
        AnalysisPlan(
            sources=[
                DatasetSource(
                    file_id=production_file_id,
                    alias="p",
                    sheet="Production",
                )
            ],
            pivot=PivotSpec(
                index=["p.Machine"],
                columns="p.Date",
                values="p.Yield",
                aggregation="mean",
            ),
        ),
        {str(production_file_id): production},
    )
    assert pivot["summary"]["result_rows"] == 2

    correlation = DatasetQueryService().execute(
        AnalysisPlan(
            sources=[
                DatasetSource(
                    file_id=production_file_id,
                    alias="p",
                    sheet="Production",
                )
            ],
            correlation=CorrelationSpec(columns=["p.Yield", "p.Pressure"]),
        ),
        {str(production_file_id): production},
    )
    assert correlation["table"]["rows"][0]["p.Yield"] == pytest.approx(1.0)


class _PlanLLM:
    async def complete(self, **_kwargs) -> str:
        return json.dumps(
            {
                "sources": [],
                "group_by": ["Machine"],
                "aggregations": [
                    {
                        "function": "mean",
                        "column": "Yield",
                        "alias": "mean_yield",
                    }
                ],
                "charts": [
                    {
                        "type": "bar",
                        "x_field": "Machine",
                        "y_field": "mean_yield",
                    }
                ],
            }
        )


@pytest.mark.asyncio
async def test_natural_language_plan_requires_explicit_confirmation(
    tmp_path: Path,
) -> None:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    principal = _principal()
    path = tmp_path / "production.xlsx"
    _workbook(path, [["Machine", "Yield"], ["A", 98], ["B", 90]])
    storage = WorkspaceStorage(root=str(tmp_path / "storage"))

    with Session(engine) as db:
        user = AuditService(db).ensure_user(principal)
        workspace = WorkspaceService(db).get_or_create_personal(principal)
        session = ChatSession(user_id=user.id, workspace_id=workspace.id)
        db.add(session)
        db.flush()
        file_id = uuid4()
        ingested = SpreadsheetIngestionService(storage).profile_and_convert(
            workspace_id=workspace.id,
            file_id=file_id,
            file_path=str(path),
            file_type="xlsx",
        )
        source = AnalysisFile(
            id=file_id,
            workspace_id=workspace.id,
            session_id=session.id,
            created_by=user.id,
            filename=path.name,
            original_filename=path.name,
            file_type="xlsx",
            file_path=str(path),
            profile_path=ingested["profile_path"],
            dataset_manifest=ingested["dataset_manifest"],
            profile_progress=100,
            status=AnalysisFileStatus.READY.value,
            size_bytes=path.stat().st_size,
            confidential_level=ConfidentialLevel.INTERNAL.value,
        )
        db.add(source)
        db.commit()

        service = AnalysisOrchestratorService(db, llm_service=_PlanLLM())
        draft = await service.create_draft(
            AnalysisPlanDraftCreate(
                workspace_id=workspace.id,
                session_id=session.id,
                question="比較每台機器的平均良率",
                file_ids=[source.id],
            ),
            principal,
        )
        assert service.response(draft).confirmation_required is True
        assert db.query(AnalysisJob).count() == 0

        job = service.confirm(draft.id, principal)
        assert job.status == "queued"
        assert job.draft_id == draft.id
        assert db.query(AnalysisJob).count() == 1


def test_report_chart_and_export_are_managed_artifacts(tmp_path: Path) -> None:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    principal = _principal()
    with Session(engine) as db:
        user = AuditService(db).ensure_user(principal)
        workspace = WorkspaceService(db).get_or_create_personal(principal)
        source = AnalysisFile(
            workspace_id=workspace.id,
            created_by=user.id,
            filename="source.csv",
            original_filename="source.csv",
            file_type="csv",
            file_path=str(tmp_path / "source.csv"),
            profile_progress=100,
            status=AnalysisFileStatus.READY.value,
            size_bytes=1,
            confidential_level=ConfidentialLevel.INTERNAL.value,
        )
        db.add(source)
        db.flush()
        job = AnalysisJob(
            workspace_id=workspace.id,
            file_id=source.id,
            created_by=user.id,
            status="completed",
            progress=100,
            request_json={"select": ["Machine", "Yield"]},
            result_json={
                "summary": {"processed_rows": 2, "matched_rows": 2, "result_rows": 2},
                "table": {
                    "columns": [
                        {"key": "Machine", "label": "Machine"},
                        {"key": "Yield", "label": "Yield"},
                    ],
                    "rows": [
                        {"Machine": "A", "Yield": 98},
                        {"Machine": "B", "Yield": 90},
                    ],
                },
                "charts": [
                    {
                        "schema_version": "2.0",
                        "type": "bar",
                        "x_field": "Machine",
                        "y_field": "Yield",
                        "data": [],
                    }
                ],
            },
        )
        db.add(job)
        db.commit()
        service = AnalysisArtifactService(
            db,
            WorkspaceStorage(root=str(tmp_path / "artifacts")),
        )
        exported = service.export(
            job.id,
            AnalysisExportRequest(format="csv"),
            principal,
        )
        report = service.create_report(
            job.id,
            AnalysisReportRequest(title="Yield report"),
            principal,
        )
        chart = service.create_chart(
            job.id,
            AnalysisChartArtifactRequest(chart_index=0),
            principal,
        )
        assert Path(exported.file_path).read_text(encoding="utf-8-sig").startswith("Machine,Yield")
        assert "# Yield report" in Path(report.file_path).read_text(encoding="utf-8")
        assert (
            json.loads(Path(chart.file_path).read_text(encoding="utf-8"))["schema_version"] == "2.0"
        )


class _HybridRetriever:
    async def retrieve(self, **_kwargs):
        return [
            RetrievedChunk(
                chunk_id=uuid4(),
                document_id=uuid4(),
                content="壓力異常時依 SOP 檢查調壓閥。",
                final_score=0.9,
                metadata={"title": "Pressure SOP", "page_start": 3, "page_end": 3},
            )
        ]


class _HybridLLM:
    async def complete(self, **kwargs) -> str:
        assert "[COMPUTED_ANALYSIS_FACTS]" in kwargs["user_prompt"]
        assert "[KNOWLEDGE_BASE_EVIDENCE]" in kwargs["user_prompt"]
        return "數據觀察引用 analysis job；處置方式依 [Source 1]。"


@pytest.mark.asyncio
async def test_hybrid_answer_combines_completed_result_and_kb_evidence() -> None:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    principal = _principal()
    with Session(engine) as db:
        user = AuditService(db).ensure_user(principal)
        workspace = WorkspaceService(db).get_or_create_personal(principal)
        source = AnalysisFile(
            workspace_id=workspace.id,
            created_by=user.id,
            filename="source.csv",
            original_filename="source.csv",
            file_type="csv",
            file_path="source.csv",
            profile_progress=100,
            status=AnalysisFileStatus.READY.value,
            size_bytes=1,
            confidential_level=ConfidentialLevel.INTERNAL.value,
        )
        db.add(source)
        db.flush()
        job = AnalysisJob(
            workspace_id=workspace.id,
            file_id=source.id,
            created_by=user.id,
            status="completed",
            progress=100,
            request_json={"select": ["Pressure"]},
            result_json={
                "summary": {"result_rows": 1},
                "table": {"columns": [], "rows": [{"Pressure": 20}]},
            },
        )
        db.add(job)
        db.commit()
        service = HybridAnalysisService(db)
        service.retriever = _HybridRetriever()
        service.llm = _HybridLLM()

        answer, citations, report = await service.answer(
            AnalysisHybridRequest(
                workspace_id=workspace.id,
                question="壓力增加後該如何處理？",
                analysis_job_ids=[job.id],
                knowledge_base_ids=[uuid4()],
                create_report=False,
            ),
            principal,
        )

        assert "Source 1" in answer
        assert citations[0].title == "Pressure SOP"
        assert report is None
